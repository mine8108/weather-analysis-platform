"""AI 报告层回归测试：视觉配置解析、多模态请求体、报告排版与 docx 嵌图。

requests 调用用打桩替换（不打真实网络），配置通过替换 st.secrets 注入。

无 pytest 时可直接运行（`python -B tests/test_ai_narrative.py`）。
"""
import io
import os
import sys

_APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _APP_DIR not in sys.path:
    sys.path.insert(0, _APP_DIR)

from PIL import Image  # noqa: E402

from modules import ai_narrative  # noqa: E402


class _FakeSecrets(dict):
    """st.secrets 的最小替身：只实现 .get。"""

    def get(self, key, default=None):
        return dict.get(self, key, default)


class _FakeResponse:
    def __init__(self, payload):
        self._payload = payload

    def raise_for_status(self):
        return None

    def json(self):
        return self._payload


def _with_secrets(values, fn):
    original = ai_narrative.st.secrets
    try:
        ai_narrative.st.secrets = _FakeSecrets(values)
        return fn()
    finally:
        ai_narrative.st.secrets = original


def _jpeg_bytes(width=64, height=48):
    buf = io.BytesIO()
    Image.new("RGB", (width, height), (20, 40, 80)).save(buf, format="JPEG")
    return buf.getvalue()


# ============================================================
# 一、视觉配置解析
# ============================================================

def test_resolve_vision_config_requires_explicit_model():
    """LLM_VISION_MODEL 未配置时视为未配置，绝不回落到 LLM_MODEL。"""
    result = _with_secrets({"LLM_API_KEY": "k", "LLM_MODEL": "deepseek-chat"},
                           ai_narrative.resolve_vision_config)
    assert result is None


def test_resolve_vision_config_returns_none_without_key():
    result = _with_secrets({"LLM_VISION_MODEL": "qwen-vl-max"},
                           ai_narrative.resolve_vision_config)
    assert result is None


def test_resolve_vision_config_falls_back_for_key_and_url():
    cfg = _with_secrets({
        "LLM_VISION_MODEL": "qwen-vl-max",
        "LLM_API_KEY": "text-key",
        "LLM_BASE_URL": "https://example.invalid/v1",
    }, ai_narrative.resolve_vision_config)
    assert cfg["api_key"] == "text-key"
    assert cfg["base_url"] == "https://example.invalid/v1"
    assert cfg["model"] == "qwen-vl-max"
    assert cfg["max_tokens"] == ai_narrative.VISION_MAX_TOKENS


def test_resolve_vision_config_prefers_vision_specific_values():
    cfg = _with_secrets({
        "LLM_VISION_MODEL": "glm-4v",
        "LLM_VISION_API_KEY": "vision-key",
        "LLM_VISION_BASE_URL": "https://vision.invalid/v1/",
        "LLM_API_KEY": "text-key",
        "LLM_BASE_URL": "https://text.invalid",
    }, ai_narrative.resolve_vision_config)
    assert cfg["api_key"] == "vision-key"
    assert cfg["base_url"] == "https://vision.invalid/v1"  # 末尾斜杠被规范化
    assert cfg["model"] == "glm-4v"


def test_resolve_vision_config_survives_missing_secrets_file():
    """secrets 文件完全缺失时 st.secrets.get 会抛异常，必须兜底为「未配置」。

    回归来源：AppTest 实跑读图页时，无 secrets 的环境直接抛
    StreamlitSecretNotFoundError，页面崩溃而不是给出配置指引。
    """
    class _RaisingSecrets:
        def get(self, key, default=None):
            raise RuntimeError("No secrets found.")

    original = ai_narrative.st.secrets
    try:
        ai_narrative.st.secrets = _RaisingSecrets()
        assert ai_narrative.resolve_vision_config() is None
    finally:
        ai_narrative.st.secrets = original


# ============================================================
# 二、多模态请求体
# ============================================================

def test_call_vision_llm_builds_multimodal_payload():
    captured = {}

    def _fake_post(url, headers=None, json=None, timeout=None):
        captured.update({"url": url, "headers": headers, "json": json,
                         "timeout": timeout})
        return _FakeResponse({"choices": [{"message": {"content": "解读正文"}}]})

    original = ai_narrative.requests.post
    try:
        ai_narrative.requests.post = _fake_post
        text = ai_narrative.call_vision_llm(
            "提示词", ["data:image/jpeg;base64,AAAA"],
            "key", base_url="https://example.invalid/v1", model="qwen-vl-max")
    finally:
        ai_narrative.requests.post = original

    assert text == "解读正文"
    assert captured["url"] == "https://example.invalid/v1/chat/completions"
    assert captured["json"]["model"] == "qwen-vl-max"
    assert captured["json"]["max_tokens"] >= 1200
    content = captured["json"]["messages"][1]["content"]
    assert content[0]["type"] == "text"
    assert content[1]["type"] == "image_url"
    assert content[1]["image_url"]["url"].startswith("data:image/jpeg;base64,")
    assert captured["headers"]["Authorization"] == "Bearer key"


def test_vision_max_tokens_is_enough_for_reasoning_models():
    """输出预算必须容得下推理型模型的思考过程。

    实测 deepseek-flash：单图 reasoning 2065 tokens，双图 10437 tokens，而六段
    正文只有约 1100 tokens。1600 与 4096 两个上限都被思考吃光，正文未开始即结束
    （finish_reason=length + content 空）。这里锁一个下界，防止再把预算改小。
    """
    assert ai_narrative.VISION_MAX_TOKENS >= 12000, ai_narrative.VISION_MAX_TOKENS


def test_call_vision_llm_accepts_max_tokens_override():
    captured = {}

    def _fake_post(url, headers=None, json=None, timeout=None):
        captured.update(json or {})
        return _FakeResponse({"choices": [{"message": {"content": "正文"}}]})

    original = ai_narrative.requests.post
    try:
        ai_narrative.requests.post = _fake_post
        ai_narrative.call_vision_llm("p", ["data:a"], "k", model="m",
                                     max_tokens=8000)
    finally:
        ai_narrative.requests.post = original
    assert captured["max_tokens"] == 8000


def test_resolve_vision_config_reads_max_tokens_override():
    """允许通过 Secrets 覆盖输出预算：合适的值取决于所选模型。"""
    class _Secrets:
        data = {"LLM_VISION_MODEL": "m", "LLM_VISION_API_KEY": "k",
                "LLM_VISION_MAX_TOKENS": "8000"}

        def get(self, key, default=None):
            return self.data.get(key, default)

    original = ai_narrative.st.secrets
    try:
        ai_narrative.st.secrets = _Secrets()
        cfg = ai_narrative.resolve_vision_config()
        assert cfg["max_tokens"] == 8000
    finally:
        ai_narrative.st.secrets = original


def test_resolve_vision_config_rejects_invalid_max_tokens():
    """非法值回落默认，不得因为一个可选键把读图功能整个弄坏。"""
    class _Secrets:
        data = {"LLM_VISION_MODEL": "m", "LLM_VISION_API_KEY": "k",
                "LLM_VISION_MAX_TOKENS": "很多"}

        def get(self, key, default=None):
            return self.data.get(key, default)

    original = ai_narrative.st.secrets
    try:
        ai_narrative.st.secrets = _Secrets()
        cfg = ai_narrative.resolve_vision_config()
        assert cfg["max_tokens"] == ai_narrative.VISION_MAX_TOKENS
    finally:
        ai_narrative.st.secrets = original


def test_call_vision_llm_supports_multiple_images():
    captured = {}

    def _fake_post(url, headers=None, json=None, timeout=None):
        captured.update(json)
        return _FakeResponse({"choices": [{"message": {"content": "x"}}]})

    original = ai_narrative.requests.post
    try:
        ai_narrative.requests.post = _fake_post
        ai_narrative.call_vision_llm("p", ["data:a", "data:b"], "k", model="m")
    finally:
        ai_narrative.requests.post = original

    images = [item for item in captured["messages"][1]["content"]
              if item["type"] == "image_url"]
    assert len(images) == 2


def _expect_failure(fn, why):
    """断言 fn 因「业务原因」失败：函数缺失或 AssertionError 都不算通过。"""
    try:
        fn()
    except AssertionError:
        raise
    except AttributeError as exc:
        raise AssertionError("目标函数不存在，测试无法成立：%s" % exc)
    except Exception:
        return
    raise AssertionError(why)


def test_call_vision_llm_raises_on_malformed_response():
    original = ai_narrative.requests.post
    try:
        ai_narrative.requests.post = lambda *a, **k: _FakeResponse({"nope": 1})
        _expect_failure(
            lambda: ai_narrative.call_vision_llm("p", ["data:a"], "k", model="m"),
            "结构异常的返回必须抛异常，不能静默返回空")
    finally:
        ai_narrative.requests.post = original


def test_call_vision_llm_raises_on_empty_content():
    original = ai_narrative.requests.post
    try:
        ai_narrative.requests.post = lambda *a, **k: _FakeResponse(
            {"choices": [{"message": {"content": "   "}}]})
        _expect_failure(
            lambda: ai_narrative.call_vision_llm("p", ["data:a"], "k", model="m"),
            "空内容必须视为失败（不生成替代文本）")
    finally:
        ai_narrative.requests.post = original


def test_call_vision_llm_requires_model():
    _expect_failure(
        lambda: ai_narrative.call_vision_llm("p", ["data:a"], "k"),
        "未指定模型时必须抛异常，不得悄悄用默认文本模型")


def _failure_message(fn):
    """执行并返回异常文本；不抛异常即判定测试不成立。"""
    try:
        fn()
    except Exception as exc:  # noqa: BLE001 - 这里就是要看异常文本
        return "%s: %s" % (exc.__class__.__name__, exc)
    raise AssertionError("预期抛异常但没有抛出")


def test_call_vision_llm_accepts_content_parts_list():
    """部分服务商把 content 返回为分片列表。旧实现只认字符串，会把正常响应误判为
    「空内容」——这正是生产环境实际遇到的失败。"""
    original = ai_narrative.requests.post
    try:
        ai_narrative.requests.post = lambda *a, **k: _FakeResponse(
            {"choices": [{"message": {"content": [
                {"type": "text", "text": "【图像信息】\nA"},
                {"type": "text", "text": "【风险提示与结论】\nB"}]},
                "finish_reason": "stop"}]})
        text = ai_narrative.call_vision_llm("p", ["data:a"], "k", model="m")
    finally:
        ai_narrative.requests.post = original
    assert "【图像信息】" in text
    assert "【风险提示与结论】" in text


def test_call_vision_llm_empty_content_reports_diagnostics():
    """空正文必须报出 finish_reason、模型名与原始响应片段。

    只说「返回了空内容」等于丢掉故障现场：上游是截断、安全拦截、还是不认图，
    这三种原因的处置完全不同。
    """
    original = ai_narrative.requests.post
    try:
        ai_narrative.requests.post = lambda *a, **k: _FakeResponse(
            {"model": "some-vl-model",
             "choices": [{"message": {"content": ""}, "finish_reason": "length"}]})
        message = _failure_message(
            lambda: ai_narrative.call_vision_llm("p", ["data:a"], "k",
                                                 model="some-vl-model"))
    finally:
        ai_narrative.requests.post = original
    assert "length" in message
    assert "some-vl-model" in message
    assert "choices" in message, message
    assert "LLM_VISION_MAX_TOKENS" in message


def test_call_vision_llm_flags_reasoning_only_response():
    """正文为空但带 reasoning_content：多为推理型模型未产出正文，须点名。"""
    original = ai_narrative.requests.post
    try:
        ai_narrative.requests.post = lambda *a, **k: _FakeResponse(
            {"choices": [{"message": {"content": "",
                                      "reasoning_content": "先看图……"},
                          "finish_reason": "stop"}]})
        message = _failure_message(
            lambda: ai_narrative.call_vision_llm("p", ["data:a"], "k", model="m"))
    finally:
        ai_narrative.requests.post = original
    assert "reasoning_content" in message


# ============================================================
# 三、报告排版与导出
# ============================================================

def test_parse_sections_splits_on_bracket_titles():
    sections = ai_narrative.parse_sections(
        "导言\n\n【图像信息】\nA\n\n【风险提示与结论】\nB\n")
    titles = [title for title, _body in sections]
    assert "图像信息" in titles
    assert "风险提示与结论" in titles


def test_build_report_meta_has_scope_and_time():
    meta = ai_narrative.build_report_meta("图片读图")
    assert meta["scope"] == "图片读图"
    assert meta["generated"]


def test_report_html_contains_title_and_sections():
    meta = ai_narrative.build_report_meta("图片读图")
    html = ai_narrative._report_html([("图像信息", "A")], meta)
    assert "图像信息" in html
    assert "<style>" in html


def test_build_docx_without_images():
    from docx import Document
    meta = ai_narrative.build_report_meta("图片读图")
    data = ai_narrative._build_docx([("图像信息", "第一段\n第二段")], meta, None)
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert any("第一段" in t for t in texts)
    assert len(doc.inline_shapes) == 0


def test_build_docx_embeds_uploaded_images():
    """读图报告必须把原图嵌进 docx，便于留档。"""
    from docx import Document
    meta = ai_narrative.build_report_meta("图片读图")
    images = [{"data_bytes": _jpeg_bytes(), "name": "a.jpg"},
              {"data_bytes": _jpeg_bytes(80, 60), "name": "b.jpg"}]
    data = ai_narrative._build_docx([("图像信息", "A")], meta, images)
    doc = Document(io.BytesIO(data))
    assert len(doc.inline_shapes) == 2


# ============================================================
# 四、已删除的旧 API 不得回流
# ============================================================

def test_detection_narrative_api_removed():
    for name in ("build_prompt", "build_fallback_markdown", "render_ai_block",
                 "call_llm", "_build_meta"):
        assert not hasattr(ai_narrative, name), name


def test_public_api_present():
    for name in ("resolve_vision_config", "call_vision_llm", "display_report",
                 "build_report_meta", "parse_sections"):
        assert hasattr(ai_narrative, name), name


if __name__ == "__main__":
    import traceback

    failed = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
            except Exception:
                failed += 1
                print(f"[FAIL] {name}")
                traceback.print_exc()
            else:
                print(f"[PASS] {name}")
    print(f"\n{'FAILED' if failed else 'ALL PASSED'} ({failed} failures)")
    sys.exit(1 if failed else 0)
