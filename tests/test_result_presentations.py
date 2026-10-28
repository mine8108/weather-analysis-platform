"""Non-network result identity, evidence, export, and parser provenance contracts."""
import io
import json
from datetime import datetime, timezone

import pandas as pd
from docx import Document

from modules.ai_report_model import SECTION_TITLES, decode_report, freeze_images, input_fingerprint, report_sections
from modules.ai_narrative import parse_sections, _build_docx, build_report_meta
from modules.codec import decode_metar, decode_synop
from modules.report.preview import document_blocks, preview_html
from modules.reporter import export_report_word


def structured():
    return {"schema_version": 1,
            "findings": [{"id": "F1", "kind": "conditional", "text": "条件性结论", "evidence_ids": ["E1"], "limitation_ids": ["L1"]}],
            "evidence": [{"id": "E1", "image_id": "img_1", "location": "图左侧", "feature": "可见 L 标记"}],
            "limitations": [{"id": "L1", "missing": "缺少时次", "impact": "无法判断演变", "next_input": "补充时次"}],
            "sections": [{"title": title, "body": "测试正文 " + title} for title in SECTION_TITLES]}


def test_snapshot_preserves_order_hash_and_cannot_rebind_to_new_uploads():
    images = [{"name": "a.png", "data_bytes": b"first", "width": 600, "height": 400, "ok": True},
              {"name": "b.png", "data_bytes": b"second", "ok": True}]
    saved = freeze_images(images)
    original = input_fingerprint(images, "note")
    images[0]["data_bytes"] = b"changed"
    assert saved[0]["data_bytes"] == b"first"
    assert [i["id"] for i in saved] == ["img_1", "img_2"]
    assert len(saved[0]["sha256"]) == 64
    assert input_fingerprint(images, "note") != original
    assert input_fingerprint(images, "new note") != input_fingerprint(images, "note")


def test_invalid_image_and_missing_limits_keep_raw_without_invented_claims():
    images = freeze_images([{"name": "a.png", "data_bytes": b"first"}])
    payload = structured()
    report = decode_report(json.dumps(payload), images)
    assert report["structured"]
    payload["evidence"][0]["image_id"] = "img_99"
    bad = decode_report(json.dumps(payload), images)
    assert not bad["structured"]
    assert not bad["findings"] and not bad["evidence"]
    assert bad["sections"][0]["title"] == "图像信息"
    payload = structured()
    payload["findings"][0]["limitation_ids"] = []
    assert not decode_report(json.dumps(payload), images)["structured"]
    payload = structured()
    payload["findings"][0]["evidence_ids"] = ["missing"]
    assert not decode_report(json.dumps(payload), images)["structured"]


def test_legacy_headings_and_inline_citations_are_readable():
    raw = "说明\n\n### 1. 【图像信息】\n内容引用【色标】不应分段。\n\n2. 【风险提示与结论】\n单图无法判断演变。"
    parsed = parse_sections(raw)
    assert [title for title, _ in parsed] == ["解读摘要", "图像信息", "风险提示与结论"]
    assert "【色标】" in parsed[1][1]
    report = decode_report(raw, [])
    assert not report["structured"] and report["raw_response"] == raw


def test_ai_word_contains_same_conclusions_evidence_limitations_and_sections():
    images = freeze_images([{"name": "a.png", "data_bytes": b"first"}])
    model = decode_report(json.dumps(structured()), images)
    sections = report_sections(model)
    data = _build_docx(sections, build_report_meta("离线单元测试"))
    texts = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert all(title in texts and body in texts for title, body in sections)
    assert "条件性结论" in texts and "可见 L 标记" in texts and "无法判断演变" in texts


def test_word_preview_reads_exact_document_in_order_and_escapes_source_text():
    frame = pd.DataFrame({"timestamp": pd.date_range("2026-10-01", periods=3, freq="h"), "temperature": [21, 22, 23]})
    for plain in (False, True):
        data = export_report_word(frame, [], 90, source="<script>alert(1)</script>", plain_language=plain)
        blocks = document_blocks(data)
        html = preview_html(data)
        assert blocks and "<script>" not in html
        assert len(Document(io.BytesIO(data)).inline_shapes) == 0
        for block in blocks:
            if block["type"] == "paragraph":
                from html import escape
                assert escape(block["text"]).replace("\n", "<br>") in html


def test_provenance_is_captured_at_actual_decoder_assignment():
    reference = datetime(2026, 10, 26, tzinfo=timezone.utc)
    metar = "METAR ZBAA 260800Z 18005MPS 9999 SCT030 24/12 Q1018 UNKNOWN"
    trace = {}
    result, errors = decode_metar(metar, reference, provenance=trace)
    assert not errors
    assert result == decode_metar(metar, reference)[0]
    assert trace["temperature"] == {"token": "24/12", "index": 6}
    assert trace["pressure"] == {"token": "Q1018", "index": 7}
    assert trace["timestamp"]["token"] == "260800Z"
    synop = "AAXX 26081 54511 11580 81805 10240 20120 30180 60001"
    trace = {}
    result, errors = decode_synop(synop, reference, provenance=trace)
    assert not errors and result == decode_synop(synop, reference)[0]
    assert trace["wind_speed"]["token"] == "81805"
    assert trace["temperature"]["token"] == "10240"


def test_unrecognized_report_has_no_fabricated_provenance():
    trace = {}
    result, errors = decode_metar("!!! NOT-A-REPORT", provenance=trace)
    assert not trace
    assert result["station_id"] == "UNKNOWN" and result["timestamp"] is None
